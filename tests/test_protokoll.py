import io

import pytest
from openpyxl import load_workbook

import config
import protokoll


def _oppna(data):
    return load_workbook(io.BytesIO(data))


@pytest.fixture(scope="module")
def fil(facit):
    rader = [dict(r, fabrikat="IMI TA", system="VS10", placering="", noteringar="") for r in facit]
    rader[1]["noteringar"] = "Temp 49 °C (pump), ca 59 °C (termometer)"
    return _oppna(protokoll.skapa_protokoll(
        rader,
        jobb={"objekt": "Testhuset", "byggnad": "B", "datum": "2026-09-20", "utfort_av": "Anna Andersson"},
        pump={"beteckning": "P1", "driftform": "Konstant tryck", "dynamiskt_mvp": "9.7", "statiskt_bar": "1,5"}))


def test_varden_hamnar_i_ratt_kolumn(fil):
    ws = fil[config.BLAD_PROTOKOLL]
    rad9 = [ws.cell(9, k).value for k in range(1, 13)]
    assert rad9 == ["VS10", 246, None, "IMI TA", "STAF*", 3.7, 65, 13.9, 24, 1.88, 1.9, None]
    assert ws.cell(18, 5).value == "STAF*" and ws.cell(18, 11).value == 3.82       # tionde och sista mätningen
    assert ws.cell(19, 5).value is None


def test_asterisken_finns_kvar_pa_alla_rader(fil):
    ws = fil[config.BLAD_PROTOKOLL]
    assert all(str(ws.cell(r, 5).value).endswith("*") for r in range(9, 19))


def test_displayens_decimaler_behalls(fil):
    ws = fil[config.BLAD_PROTOKOLL]
    assert ws["I9"].number_format == "0.0"      # 24.0
    assert ws["K9"].number_format == "0.00"     # 1.90
    assert ws["G9"].number_format == "0"        # 65
    assert ws["K14"].number_format == "0.000"   # 0.021


def test_mallens_utseende_ar_kvar(fil):
    ws = fil[config.BLAD_PROTOKOLL]
    assert len(ws._images) == 1 and len(fil[config.BLAD_FORSATTSBLAD]._images) == 1    # loggan
    kant = ws["I12"].border
    assert (kant.left.style, kant.top.style) == ("thin", "hair")                       # rutnätet
    assert ws["I12"].font.name == "Arial" and ws["I12"].alignment.horizontal == "center"
    assert ws.print_title_rows == "$1:$8" and ws.freeze_panes == "A9"
    assert ws["F6"].value == "Inställning" and ws["G4"].value == "l/s"                 # rubrikerna orörda


def test_notering_far_radbrytning(fil):
    ws = fil[config.BLAD_PROTOKOLL]
    assert ws["L10"].value.startswith("Temp 49") and ws["L10"].alignment.wrap_text
    assert ws.row_dimensions[10].height == 24


def test_rubrikfalt_pa_bada_bladen(fil):
    ws, fs = fil[config.BLAD_PROTOKOLL], fil[config.BLAD_FORSATTSBLAD]
    assert (ws["G2"].value, ws["G3"].value, ws["L2"].value, ws["A5"].value) == ("Testhuset", "B", "2026-09-20", "Anna Andersson")
    assert (fs["J2"].value, fs["J3"].value, fs["O2"].value, fs["A4"].value) == ("Testhuset", "B", "2026-09-20", "Anna Andersson")


def test_mallen_delas_av_flera_inget_far_folja_med_fran_forra_jobbet():
    """Mallen kom med 'Söderhallarna' och ett namn förifyllt. Det får aldrig hamna i en kollegas protokoll."""
    mall = load_workbook(config.MALL if False else __import__("pathlib").Path(protokoll.__file__).with_name(config.MALL))
    for blad in mall.worksheets:
        text = " ".join(str(c.value) for rad in blad.iter_rows() for c in rad if c.value)
        assert "Söderhallarna" not in text and "Hayrapetian" not in text and "VS10" not in text
    wb = _oppna(protokoll.skapa_protokoll([{"typ": "STAD*"}], jobb={"objekt": "Kv. Eken", "byggnad": "", "datum": "2026-09-21", "utfort_av": "Kalle K"}))
    assert wb[config.BLAD_PROTOKOLL]["G3"].value == "-" and wb[config.BLAD_FORSATTSBLAD]["J3"].value == "-"     # tom byggnad visas som i mallen
    assert wb[config.BLAD_PROTOKOLL]["A5"].value == "Kalle K"


def test_pumpdata_pa_forsattsbladet(fil):
    fs = fil[config.BLAD_FORSATTSBLAD]
    assert fs["A10"].value == "Pump P1 inställd på:"
    assert fs["A13"].value == "Driftform: Konstant tryck"
    assert fs["A14"].value == "Dynamisk tryck: 9,7 mvp (95 kPa)"
    assert fs["A15"].value == "Statisk tryck: 1,5 Bar"


def test_utan_pump_och_jobb_lamnas_mallen_orord():
    wb = _oppna(protokoll.skapa_protokoll([{"typ": "STAD*", "uppmatt": "0.192"}]))
    fs, ws = wb[config.BLAD_FORSATTSBLAD], wb[config.BLAD_PROTOKOLL]
    assert fs["A14"].value == "Dynamisk tryck:  mvp ( kPa)" and ws["G2"].value is None
    assert ws["A9"].value is None and ws["A10"].value is None
    assert ws["K9"].value == 0.192


def test_text_i_talkolumn_och_ventilnummer_med_bokstaver():
    wb = _oppna(protokoll.skapa_protokoll([{"dimension": "DN20", "ventilnummer": "RV-07", "kv": "2,7"}]))
    ws = wb[config.BLAD_PROTOKOLL]
    assert (ws["G9"].value, ws["B9"].value, ws["H9"].value) == ("DN20", "RV-07", 2.7)
