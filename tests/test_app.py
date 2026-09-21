"""Röktest av hela appen i Streamlits testläge – utan riktiga API-anrop."""
import io

import pytest
from conftest import ROT, display, svar
from openpyxl import load_workbook
from PIL import Image
from streamlit.testing.v1 import AppTest

import avlasning


def _jpeg():
    ut = io.BytesIO()
    Image.new("RGB", (300, 400), "white").save(ut, "JPEG")
    return ut.getvalue()


def _post(seq, namn, s, skillnader=None):
    rader = avlasning.till_rader(s, namn)
    return {"namn": namn, "seq": seq, "fel": None, "jpeg": _jpeg(), "rader": rader, "skillnader": skillnader or {},
            "olika_antal": False, "pump": avlasning.pumpforslag(s), "bildkvalitet": s["bildkvalitet"],
            **avlasning.sorteringsinfo(namn, b"")}


def _md(app):
    return "\n".join(m.value for m in app.markdown)


@pytest.fixture
def app(monkeypatch):
    monkeypatch.setattr("sakerhet.FORDROJNING_VID_FEL", 0)          # testerna ska inte vänta på fördröjningen
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test")
    monkeypatch.delenv("APP_LOSENORD", raising=False)
    monkeypatch.delenv("SMTP_ANVANDARE", raising=False)
    monkeypatch.delenv("SMTP_LOSENORD", raising=False)
    monkeypatch.chdir(ROT)
    return AppTest.from_file(str(ROT / "app.py"), default_timeout=30)


def test_tom_app_startar(app):
    app.run()
    assert not app.exception
    assert [b.label for b in app.button] == ["Läs av bilderna", "Prova med exempelbilder"] and app.button[0].disabled
    assert "Ladda upp dina mätningar" in "".join(m.value for m in app.markdown)


def test_demoknappen_laddar_tio_matningar_utan_api(app):
    app.run()
    app.button[1].click().run()
    assert not app.exception
    assert "<b>10</b><span>mätningar</span>" in _md(app) and "<b>0</b><span>går inte ihop</span>" in _md(app) \
        and "<b>2</b><span>osäkra</span>" in _md(app)
    knappar = {b.label: b for b in app.button}
    assert "Lägg till fler bilder" in knappar and knappar["Fortsätt till export"].disabled
    assert any("Kryssa i rutan ovan" in c.value for c in app.caption)               # och appen SÄGER varför knappen är släckt


def test_losenord_stoppar_obehoriga(app, monkeypatch):
    monkeypatch.setenv("APP_LOSENORD", "hemligt")
    app.run()
    assert not app.exception and len(app.text_input) == 1
    assert [b.label for b in app.button] == ["Logga in"]                            # bara inloggningsrutan syns
    app.text_input[0].set_value("fel")
    app.button[0].click().run()
    assert "Fel lösenord." in [e.value for e in app.error]
    app.text_input[0].set_value("hemligt")
    app.button[0].click().run()
    assert not app.exception and "Läs av bilderna" in [b.label for b in app.button]   # rätt lösenord släpper in
    assert any("Lösenordet till appen är svagt" in w.value for w in app.warning)     # ...men "hemligt" får en varning


def test_starkt_losenord_ger_ingen_varning(app, monkeypatch):
    monkeypatch.setenv("APP_LOSENORD", "ventil-gurka-tåg-fjorton")
    app.run()
    app.text_input[0].set_value("ventil-gurka-tåg-fjorton")
    app.button[0].click().run()
    assert not app.exception and not app.warning


def _laddad(app):
    ren = _post(2, "IMG_2033.jpeg", svar())
    trasig = _post(1, "IMG_2031.jpeg", svar([display(kv="2.1", klocka="06:29", osakra_falt=["installning"])]),      # Kv går inte ihop
                   skillnader={0: [("dimension", "25", "26")]})
    pump = _post(3, "IMG_2040.jpeg", svar(
        [display(flode="0.741", dp="3.51", projekterat_flode="0", procent="-.-", ventil="STAD* 32", installning="4", kv="14.2")],
        pumpar=[{"modell": "Grundfos MAGNA3", "reglertyp": "Konst. tryck", "driftsform": "Max.", "uppfordringshojd": "9.7",
                 "uppfordringshojd_enhet": "m", "flode": "", "flode_enhet": "", "vatsketemperatur": "49", "osaker": False}]))
    misslyckad = {"namn": "IMG_2050.jpeg", "seq": 4, "fel": "Fick ingen kontakt med AI-tjänsten.", "rader": [],
                  "skillnader": {}, "olika_antal": False, "pump": None, "bildkvalitet": "", **avlasning.sorteringsinfo("IMG_2050.jpeg", b"")}
    app.session_state["resultat"] = {"a": ren, "b": trasig, "c": pump, "d": misslyckad}
    app.session_state["seq"] = 4
    app.session_state["steg"] = 2
    return app.run()


def _knapp(app, etikett):
    return next(b for b in app.button if b.label == etikett)


def test_steg_2_granskning_visar_flaggor_och_en_tydlig_vag_vidare(app):
    _laddad(app)
    assert not app.exception
    text = "\n".join(m.value for m in app.markdown)
    assert text.index("Mätning 1") < text.index("Mätning 2")                        # sorterat efter filnummer: 2031 före 2033
    assert "Kv går inte ihop" in text and "AI:n är osäker på inställning" in text
    assert "de två avläsningarna skiljer sig (&#x27;25&#x27; / &#x27;26&#x27;)" in text       # citattecknen är HTML-säkrade
    assert text.count(">Stämmer<") == 2 and text.count(">Kontrollera<") >= 1      # ren + pumpbilden (bara en upplysning)
    assert "Inget projekterat flöde" in text and "Klar med granskningen?" in text
    assert any("IMG_2050.jpeg" in e.value for e in app.error)
    assert any("kunde inte läsas av" in w.value for w in app.warning)
    assert "<b>3</b><span>mätningar</span>" in text and "<b>1</b><span>går inte ihop</span>" in text and "<b>2</b><span>osäkra</span>" in text
    assert _knapp(app, "Fortsätt till export").disabled                             # släckt tills flaggorna är kvitterade ...
    app.checkbox[-1].check().run()
    assert not _knapp(app, "Fortsätt till export").disabled                         # ... sedan tänd


def test_steg_3_export_hela_vagen_till_nedladdning(app):
    _laddad(app)
    app.checkbox[-1].check().run()
    _knapp(app, "Fortsätt till export").click().run()
    assert not app.exception and app.session_state["steg"] == 3
    falt = {t.label: t for t in app.text_input}
    assert falt["Dynamiskt tryck (mvp)"].value == "9.7" and falt["Driftform"].value == "Konstant tryck"
    assert not app.get("download_button")

    _knapp(app, "Skapa protokollet").click().run()                                  # utan obligatoriska fält: tydligt besked
    assert any("Fyll i Objekt och Utfört av" in e.value for e in app.error) and not app.get("download_button")

    falt = {t.label: t for t in app.text_input}
    falt["Objekt *"].set_value("Kv. Eken")
    falt["Utfört av *"].set_value("Kalle K")
    falt["System"].set_value("VS10")
    _knapp(app, "Skapa protokollet").click().run()
    assert not app.exception and len(app.get("download_button")) == 1
    assert any("Injusteringsprotokoll_Kv._Eken_" in x.value for x in app.success)

    wb = load_workbook(io.BytesIO(app.session_state["fil"]["data"]))                 # och filen innehåller det den ska
    ws = wb["Injust.protokoll"]
    assert (ws["G2"].value, ws["A5"].value, ws["A9"].value, ws["E9"].value, ws["K9"].value) == ("Kv. Eken", "Kalle K", "VS10", "STAD*", 0.192)
    assert ws["E11"].value == "STAD*" and ws["J11"].value is None and ws["E12"].value is None      # 3 mätningar, inget projekterat på den tredje
    assert wb["Försättsblad"]["A14"].value == "Dynamisk tryck: 9,7 mvp (95 kPa)"


def test_tillbaka_och_nytt_jobb(app):
    _laddad(app)
    app.checkbox[-1].check().run()
    _knapp(app, "Fortsätt till export").click().run()
    _knapp(app, "‹ Tillbaka till granskningen").click().run()
    assert not app.exception and app.session_state["steg"] == 2
    assert app.checkbox[-1].value is True                                            # kvitteringen finns kvar
    _knapp(app, "Fortsätt till export").click().run()
    falt = {t.label: t for t in app.text_input}
    falt["Objekt *"].set_value("Kv. Eken")
    falt["Utfört av *"].set_value("Kalle K")
    _knapp(app, "Skapa protokollet").click().run()
    _knapp(app, "Börja på ett nytt jobb").click().run()
    assert not app.exception and app.session_state["steg"] == 1 and not app.session_state["resultat"]
    assert app.session_state["jobb"] == {"utfort_av": "Kalle K", "mejl": ""}         # namnet följer med till nästa jobb


def test_protokollet_mejlas_nar_adress_ar_ifylld(app, monkeypatch):
    import mejl
    skickat = {}
    monkeypatch.setenv("SMTP_ANVANDARE", "appen@gmail.com")
    monkeypatch.setenv("SMTP_LOSENORD", "abcd efgh ijkl mnop")
    monkeypatch.setattr(mejl, "skicka", lambda inst, till, filnamn, data, *resten: skickat.update(till=till, filnamn=filnamn, storlek=len(data), resten=resten))
    _laddad(app)
    app.checkbox[-1].check().run()
    _knapp(app, "Fortsätt till export").click().run()
    falt = {t.label: t for t in app.text_input}
    falt["Objekt *"].set_value("Kv. Eken")
    falt["Utfört av *"].set_value("Kalle K")
    falt["Mejla protokollet till"].set_value("kalle@foretaget.se")
    _knapp(app, "Skapa protokollet").click().run()
    assert not app.exception
    assert skickat["till"] == "kalle@foretaget.se" and skickat["filnamn"].startswith("Injusteringsprotokoll_Kv._Eken_") and skickat["storlek"] > 10_000
    assert skickat["resten"][0] == "Kv. Eken" and skickat["resten"][2] == 3 and skickat["resten"][3] == "Kalle K"
    assert any("skickat till **kalle@foretaget.se**" in x.value for x in app.success)
    assert len(app.get("download_button")) == 1                                     # nedladdning finns alltid kvar
    _knapp(app, "Börja på ett nytt jobb").click().run()
    assert app.session_state["jobb"] == {"utfort_av": "Kalle K", "mejl": "kalle@foretaget.se"}   # adressen följer med


def test_mejlfel_stoppar_inte_protokollet(app, monkeypatch):
    import mejl

    def trasigt(*a, **kw):
        raise mejl.MejlFel("Ingen kontakt med mejlservern.")
    monkeypatch.setenv("SMTP_ANVANDARE", "appen@gmail.com")
    monkeypatch.setenv("SMTP_LOSENORD", "abcd efgh ijkl mnop")
    monkeypatch.setattr(mejl, "skicka", trasigt)
    _laddad(app)
    app.checkbox[-1].check().run()
    _knapp(app, "Fortsätt till export").click().run()
    falt = {t.label: t for t in app.text_input}
    falt["Objekt *"].set_value("Kv. Eken"); falt["Utfört av *"].set_value("Kalle K"); falt["Mejla protokollet till"].set_value("kalle@foretaget.se")
    _knapp(app, "Skapa protokollet").click().run()
    assert any("mejlet gick inte iväg" in e.value for e in app.error) and len(app.get("download_button")) == 1


def test_utan_mejlinstallning_syns_faltet_anda_grat_med_forklaring(app):
    """Det som hände på riktigt: fältet var helt dolt, och då 'gick det inte att välja mejl'."""
    _laddad(app)
    app.checkbox[-1].check().run()
    _knapp(app, "Fortsätt till export").click().run()
    falt = {t.label: t for t in app.text_input}
    assert falt["Mejla protokollet till"].disabled and "Inte påslaget" in falt["Mejla protokollet till"].placeholder
    assert any("Varför är fältet grått?" in e.label for e in app.expander)
    assert "SMTP_ANVANDARE" in _md(app) and "fungerar i alla webbläsare" in _md(app)


def test_ifyllda_men_bortkommenterade_mejlrader_upptacks(app, tmp_path, monkeypatch):
    import hemligheter
    (tmp_path / ".streamlit").mkdir()
    (tmp_path / ".streamlit" / "secrets.toml").write_text(
        'ANTHROPIC_API_KEY = "x"\n# SMTP_ANVANDARE = "kvittens.he@gmail.com"\n# SMTP_LOSENORD = "abcd efgh ijkl mnop"\n'
        '# SMTP_SERVER = "..."\n', encoding="utf8")
    assert hemligheter.bortkommenterade(tmp_path) == ["SMTP_ANVANDARE", "SMTP_LOSENORD"]
    (tmp_path / ".streamlit" / "secrets.toml").write_text(hemligheter.MALL, encoding="utf8")
    assert hemligheter.bortkommenterade(tmp_path) == []                              # mallens exempelrader räknas inte


def test_efter_fem_felaktiga_losenord_sparras_inloggningen(app, monkeypatch):
    import lagring
    lagring._lager().pop("sparr", None)
    monkeypatch.setenv("APP_LOSENORD", "ventil-gurka-tåg-fjorton")
    app.run()
    for _ in range(5):
        app.text_input[0].set_value("gissning")
        app.button[0].click().run()
    assert not app.exception and any("spärrad" in e.value for e in app.error)
    assert not app.text_input                                                        # inget fält att gissa i längre
    lagring._lager().pop("sparr", None)


def test_samma_ventilnummer_varnar_och_kraver_kvittering(app):
    a_ = _post(1, "IMG_1.jpeg", svar([display(ventilnummer="246")]))
    b_ = _post(2, "IMG_2.jpeg", svar([display(ventilnummer="246", klocka="07:00", flode="0.238", dp="6.49", kv="3.36",
                                                 projekterat_flode="0.24", procent="99", ventil="STAD* 32", installning="1.5")]))
    app.session_state["resultat"] = {"a": a_, "b": b_}
    app.session_state["seq"] = 2
    app.session_state["steg"] = 2
    app.run()
    assert not app.exception and any("Ventilnummer **246** står på flera mätningar" in w.value for w in app.warning)
    assert _knapp(app, "Fortsätt till export").disabled                             # måste kvitteras


def test_integritetsrutan_och_radera_jobbet(app):
    _laddad(app)
    assert any("Om dina bilder och uppgifter" in e.label for e in app.expander)
    assert "Inget sparas på disk" in _md(app)
    _knapp(app, "Radera det här jobbet nu").click().run()
    assert not app.exception and not app.session_state["resultat"] and app.session_state["steg"] == 1


def test_flera_pumpar_gar_att_valja(app):
    pump = {"modell": "Grundfos MAGNA3", "reglertyp": "Konst. tryck", "driftsform": "Auto", "uppfordringshojd": "9.7",
            "uppfordringshojd_enhet": "m", "flode": "", "flode_enhet": "", "vatsketemperatur": "", "osaker": False}
    s2 = svar(pumpar=[pump, dict(pump, modell="Wilo Stratos", uppfordringshojd="4.5")])
    post = _post(1, "IMG_1.jpeg", s2)
    post["pumpar"] = avlasning.alla_pumpforslag(s2)
    app.session_state["resultat"] = {"a": post}
    app.session_state["seq"] = 1
    app.session_state["steg"] = 2
    app.run()
    _knapp(app, "Fortsätt till export").click().run()
    assert not app.exception and len(app.selectbox) == 1
    assert {t.label: t.value for t in app.text_input}["Dynamiskt tryck (mvp)"] == "9.7"
    app.selectbox[0].select(1).run()
    assert {t.label: t.value for t in app.text_input}["Dynamiskt tryck (mvp)"] == "4.5"
